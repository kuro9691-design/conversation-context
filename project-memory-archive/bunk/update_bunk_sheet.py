#!/usr/bin/env python3
"""
BUNK 投稿カレンダー（xlsx）への行反映スクリプト

新しい巻のプロンプト（generate_bunk_prompts.py）とキャプション
（build_bunk_caption.py）が確定したあと、その内容を投稿カレンダーの
xlsx に1行追加する。書式（フォント・罫線・折り返し・行の高さ）は
Vol.1の行（2行目）からそのままコピーする。P列の文字数だけは
式 =LEN(J{row}) を書く。

前提: 行番号 = 巻数 + 1（2行目=Vol.1, 3行目=Vol.2, 4行目=Vol.3 ...）

使い方:
  1. row_config_example.json を複製して埋める
  2. python update_bunk_sheet.py row_volN.json
  3. 書き込み後、必ずLibreOfficeベースのrecalcで再計算する
"""

import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

COLUMN_COUNT = 16


def main():
    if len(sys.argv) != 2:
        print("使い方: python update_bunk_sheet.py row_volN.json")
        sys.exit(1)

    config_path = Path(sys.argv[1])
    cfg = json.loads(config_path.read_text(encoding="utf-8"))

    xlsx_path = Path(cfg["xlsx_path"])
    wb = load_workbook(xlsx_path)
    ws = wb["BUNK 144回"] if "BUNK 144回" in wb.sheetnames else wb.active

    volume = cfg["volume"]
    row = volume + 1

    values = [
        volume,
        cfg["post_date"],
        cfg["month"],
        cfg["era_jp"],
        cfg["era_en"],
        cfg["period"],
        cfg["category"],
        cfg["title_jp"],
        cfg["title_en"],
        cfg["full_caption"],
        cfg["jp_body"],
        cfg["en_body"],
        cfg["hashtags"],
        cfg["body_prompts"],
        cfg["cover_prompt"],
        None,
    ]

    for col, val in enumerate(values, start=1):
        ws.cell(row=row, column=col, value=val)

    p_col = get_column_letter(16)
    ws[f"{p_col}{row}"] = f"=LEN(J{row})"

    src_row = 2
    for col in range(1, COLUMN_COUNT + 1):
        src_cell = ws.cell(row=src_row, column=col)
        dst_cell = ws.cell(row=row, column=col)
        if src_cell.font:
            dst_cell.font = src_cell.font.copy()
        if src_cell.alignment:
            dst_cell.alignment = src_cell.alignment.copy()
        if src_cell.border:
            dst_cell.border = src_cell.border.copy()

    if src_row in ws.row_dimensions and ws.row_dimensions[src_row].height:
        ws.row_dimensions[row].height = ws.row_dimensions[src_row].height

    wb.save(xlsx_path)
    print(f"{xlsx_path.name} の {row}行目（Vol.{volume}）に書き込み完了。")
    print("次に recalc で再計算してください（P列のLEN式に値を入れるため）。")


if __name__ == "__main__":
    main()
